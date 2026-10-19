"""Document text extraction: .txt, .docx, PDF (with per-page OCR for scanned pages) and images."""
from __future__ import annotations

import os
from dataclasses import dataclass

IMAGE_EXT = (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".heic")
SUPPORTED_EXT = (".txt", ".pdf", ".docx") + IMAGE_EXT
FILE_FILTER = "Documents and images (*.txt *.pdf *.docx *.png *.jpg *.jpeg *.tif *.tiff *.heic)"
MIN_PAGE_TEXT = 20          # a PDF page with fewer characters than this is treated as a scan


class DocumentError(Exception):
    pass


@dataclass
class LoadedDoc:
    text: str
    pages: int = 0
    ocr_pages: int = 0
    engine: str = ""
    renderer: str = ""


def _read_txt(path: str) -> str:
    with open(path, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _read_docx(path: str) -> str:
    import docx
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append("\t".join(c.text for c in row.cells))
    return "\n".join(parts)


def _need_ocr(what: str):
    return DocumentError(f"{what} needs OCR. Turn on “Read scanned pages and images” in Settings > Documents, "
                         "and check that an OCR engine is available there.")


def _read_pdf(path: str, ocr, languages, progress) -> LoadedDoc:
    import pypdf
    reader = pypdf.PdfReader(path)
    if reader.is_encrypted:
        raise DocumentError("This PDF is password-protected. Remove the password first.")
    texts = [(page.extract_text() or "").strip() for page in reader.pages]
    scanned = [i for i, t in enumerate(texts) if len(t) < MIN_PAGE_TEXT]
    ocr_done, used = 0, ""
    if scanned and ocr is not None:
        from aegis import pdfrender
        try:
            with pdfrender.PdfRenderer(path) as renderer:
                for n, i in enumerate(scanned):
                    if progress:
                        progress(n, len(scanned))
                    texts[i] = ocr.recognize(renderer.render(i), languages).strip()
                    ocr_done += 1
                used = renderer.used
        except pdfrender.RenderError as e:
            raise DocumentError(f"This looks like a scanned PDF, but its pages could not be turned into images. {e}") from None
    text = "\n\n".join(t for t in texts if t)
    if not text:
        raise _need_ocr("This looks like a scanned PDF, which") if ocr is None else DocumentError(
            "No text could be read from this PDF, even with OCR.")
    return LoadedDoc(text, len(texts), ocr_done, getattr(ocr, "name", "") if ocr_done else "", used)


def _read_image(path: str, ocr, languages) -> LoadedDoc:
    if ocr is None:
        raise _need_ocr("An image")
    from aegis import ocr as ocrmod
    try:
        from PIL import Image
        png = ocrmod.image_to_png_bytes(Image.open(path))
    except Exception:
        with open(path, "rb") as f:
            png = f.read()      # e.g. HEIC: Apple Vision can read the original bytes
    text = ocr.recognize(png, languages).strip()
    if not text:
        raise DocumentError("No text could be read from this image.")
    return LoadedDoc(text, 1, 1, ocr.name)


def load_document(path: str, ocr=None, languages=("en",), progress=None) -> LoadedDoc:
    ext = os.path.splitext(path)[1].lower()
    if ext == ".txt":
        return LoadedDoc(_read_txt(path))
    if ext == ".docx":
        return LoadedDoc(_read_docx(path))
    if ext == ".pdf":
        return _read_pdf(path, ocr, languages, progress)
    if ext in IMAGE_EXT:
        return _read_image(path, ocr, languages)
    raise DocumentError(f"Unsupported file type: {ext or 'unknown'}")


def load_text(path: str) -> str:
    """Compatibility helper: text only, no OCR."""
    return load_document(path).text
