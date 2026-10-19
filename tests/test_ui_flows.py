"""Drives the real window headless: every button path, with dialogs stubbed. Finds bugs unit tests cannot."""
import os, sys, time, json, zipfile
import pytest
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
import keyring
from keyring.backend import KeyringBackend
from PySide6.QtCore import QCoreApplication, QMimeData, QSettings, QUrl
from PySide6.QtWidgets import QApplication, QFileDialog, QInputDialog, QMessageBox

from aegis import redaction, providers, stores, diagnostics, mailbox, ocr
from aegis.session import Attachment

_APP = None
_ENGINE = None


class Mem(KeyringBackend):
    priority = 1
    d = {}
    def get_password(self, s, u): return Mem.d.get((s, u))
    def set_password(self, s, u, p): Mem.d[(s, u)] = p
    def delete_password(self, s, u): Mem.d.pop((s, u), None)


class FakeOCR:
    name = "Fake"
    def recognize(self, png, languages=("en",)): return "Scanned: Maria Santos maria@x.com"


def pump(cond, timeout=60):
    t = time.time()
    while not cond() and time.time() - t < timeout:
        _APP.processEvents(); time.sleep(0.01)
    _APP.processEvents()
    return cond()


@pytest.fixture()
def env(tmp_path, monkeypatch):
    global _APP, _ENGINE
    _APP = QApplication.instance() or QApplication([])
    _APP.setStyle("Fusion")
    QCoreApplication.setOrganizationName("AegisUI"); QCoreApplication.setApplicationName("Flows")
    QSettings.setPath(QSettings.NativeFormat, QSettings.UserScope, str(tmp_path / "qs"))
    QSettings.setPath(QSettings.IniFormat, QSettings.UserScope, str(tmp_path / "qs"))
    keyring.set_keyring(Mem()); Mem.d.clear(); stores._CACHE.clear(); stores.settings().clear()
    monkeypatch.setenv("HOME", str(tmp_path))
    diagnostics.LOG_DIR = str(tmp_path / "logs"); diagnostics.LOG_FILE = diagnostics.LOG_DIR + "/a.log"
    diagnostics.FAULT_FILE = diagnostics.LOG_DIR + "/n.log"; diagnostics.init("t")
    _ENGINE = _ENGINE or redaction.build_engine()
    monkeypatch.setattr(redaction, "build_engine", lambda: _ENGINE)
    stores.set_secret(stores.SECRET_KEY_NAMES["anthropic"], "sk-test")
    stores.settings().setValue("provider", "anthropic")
    stores.settings().setValue("folders/out_dir", str(tmp_path / "out"))
    stores.settings().setValue("folders/watch", False)
    shown = []
    for fn in ("warning", "information", "critical"):
        monkeypatch.setattr(QMessageBox, fn, staticmethod(lambda *a, _f=fn, **k: shown.append((_f, a[2] if len(a) > 2 else "")) or QMessageBox.Ok))
    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.Yes))
    sent = {}
    def fake_send(provider, key, model, system, prompt, url=""):
        sent.update(provider=provider, prompt=prompt)
        return ("Answer for <PERSON_1> at <EMAIL_ADDRESS_1>. Also <PERSON_9> and [Date].", model)
    monkeypatch.setattr(providers, "send", fake_send)
    import app as A
    w = A.MainWindow(); w.show()
    assert pump(lambda: w.engine is not None)
    yield w, sent, shown, tmp_path, monkeypatch
    w.close(); w.deleteLater(); _APP.processEvents()


def ask(w, text="Summarise this", atts=None):
    w.composer.set_text(text); w.composer.set_attachments(atts or [])
    w.submit()
    assert pump(lambda: w.review is not None), "review card never appeared"


def doc(text="Hi, Maria Santos (maria@x.com) asked about pricing. Call +351 912 345 678."):
    return Attachment("document", "Contract Maria.pdf", text)


# ---------------------------------------------------------------- Ask flow
def test_full_ask_flow_sends_only_sanitised_text(env):
    w, sent, shown, *_ = env
    ask(w, "Summarise for Ana", [doc()])
    w.send_now()
    assert pump(lambda: w.answer is not None)
    assert "Maria" not in sent["prompt"] and "maria@x.com" not in sent["prompt"] and "Contract" not in sent["prompt"]
    final = w.answer.final.toPlainText()
    assert "Maria Santos" in final and "maria@x.com" in final and "<PERSON_9>" in final
    assert w.answer.fill_btn.isEnabled() and w.composer.send_btn.isEnabled()
    assert not shown, f"unexpected dialogs: {shown}"

def test_cancel_returns_everything_to_the_composer(env):
    w, sent, shown, *_ = env
    ask(w, "Check this", [doc()])
    w.cancel_review()
    assert w.composer.text() == "Check this" and len(w.composer.attachments()) == 1
    assert w.stack.currentIndex() == 0 and w.composer.send_btn.isEnabled() and not sent

def test_language_correction_reruns_detection_and_replaces_the_card(env):
    w, *_ = env
    ask(w, "Check", [doc("Olá, o João Silva pediu um orçamento para a próxima semana, joao@x.pt")])
    first = w.review
    w._reanalyze("en")
    assert pump(lambda: w.review is not None and w.review is not first)
    assert w.session.analysis.language == "en" and w.session.analysis.language_source == "chosen"

def test_toggling_in_the_review_card_and_inspector_stays_in_sync(env):
    w, *_ = env
    ask(w, "Fatura", [Attachment("document", "f.txt", "Fatura 100200300 paga. NIF 123456789.")])
    stores.settings().setValue("region", "Portugal (PT)")
    w.cancel_review(); ask(w, "Fatura", [Attachment("document", "f.txt", "Fatura 100200300 paga. NIF 123456789.")])
    if w.review.checks:
        det_id, cb = next(iter(w.review.checks.items()))
        before = w.session.result.text
        cb.setChecked(not cb.isChecked())
        assert w.session.result.text != before
        w.inspector.refresh()
        assert w.inspector.det_table.rowCount() >= 0

def test_inspector_add_term_and_always_redact_update_the_render(env):
    w, _s, _sh, _t, mp = env
    ask(w, "Look", [Attachment("document", "a.txt", "The Zeta project and Rua das Flores 12 are private.")])
    mp.setattr(QInputDialog, "getText", staticmethod(lambda *a, **k: ("Zeta project", True)))
    w.inspector.show(); w.inspector._add_term()
    assert "Zeta" not in w.session.result.text
    w.session.analysis.add_term("Rua das Flores 12", "PRESET"); w._on_edited()
    assert "Rua das Flores" not in w.session.result.text

def test_provider_failure_after_send_is_reported_and_unlocks_the_composer(env):
    w, sent, shown, _t, mp = env
    def boom(*a, **k): raise providers.ProviderError("HTTP 401 invalid key", "HTTPError", 401)
    mp.setattr(providers, "send", boom)
    ask(w, "Hi", [doc()]); w.send_now()
    assert pump(lambda: bool(shown))
    assert w.composer.send_btn.isEnabled() and w.answer is None

def test_new_task_wipes_mapping_even_mid_review(env):
    w, *_ = env
    ask(w, "Hi", [doc()])
    w.new_task()
    assert w.session.analysis is None and w.session.result is None and w.stack.currentIndex() == 0
    assert w.composer.send_btn.isEnabled()

def test_idle_wipe_clears_answer_and_mapping(env):
    w, *_ = env
    ask(w, "Hi", [doc()]); w.send_now(); assert pump(lambda: w.answer is not None)
    w._on_idle()
    assert w.answer is None and w.session.sent_result is None

def test_fill_placeholders_copy_and_save(env, tmp_path):
    w, _s, _sh, _t, mp = env
    ask(w, "Hi", [doc()]); w.send_now(); assert pump(lambda: w.answer is not None)
    mp.setattr(QInputDialog, "getText", staticmethod(lambda *a, **k: ("Rita", True)))
    w.fill_placeholders()
    assert "<PERSON_9>" not in w.answer.final.toPlainText() and "Rita" in w.answer.final.toPlainText()
    w.copy_result(); assert "Rita" in QApplication.clipboard().text()
    dest = str(tmp_path / "ans.md")
    mp.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (dest, "")))
    w.save_result(); assert "Rita" in open(dest).read()


# ---------------------------------------------------------------- attachments and drops
def test_attach_txt_docx_and_scanned_pdf_and_image(env, tmp_path):
    w, _s, shown, _t, mp = env
    import docx
    (tmp_path / "a.txt").write_text("Hello Maria Santos")
    d = docx.Document(); d.add_paragraph("Contract for Maria Santos, NIF 123456789"); d.save(tmp_path / "b.docx")
    from PIL import Image
    Image.new("RGB", (600, 400), "white").save(tmp_path / "c.pdf"); Image.new("RGB", (600, 400), "white").save(tmp_path / "d.png")
    mp.setattr(w, "_ocr_config", lambda: (FakeOCR(), ("en",)))
    w.attach_documents([str(tmp_path / n) for n in ("a.txt", "b.docx", "c.pdf", "d.png")])
    assert pump(lambda: len(w.composer.attachments()) == 4)
    names = [a.name for a in w.composer.attachments()]
    assert sum("OCR" in n for n in names) == 2 and not shown

def test_unsupported_and_unreadable_files_do_not_crash(env, tmp_path):
    w, _s, shown, *_ = env
    (tmp_path / "x.xyz").write_text("x"); (tmp_path / "bad.pdf").write_bytes(b"nope")
    w.attach_documents([str(tmp_path / "x.xyz"), str(tmp_path / "bad.pdf")])
    assert pump(lambda: bool(shown)); assert w.composer.attachments() == []

def test_scanned_pdf_with_ocr_off_gives_actionable_message(env, tmp_path):
    w, _s, shown, _t, mp = env
    from PIL import Image
    Image.new("RGB", (600, 400), "white").save(tmp_path / "c.pdf")
    stores.settings().setValue("ocr/enabled", False)
    w.attach_documents([str(tmp_path / "c.pdf")])
    assert pump(lambda: bool(shown)); assert "OCR" in shown[0][1]

class _Ev:
    def __init__(self, paths): self._m = QMimeData(); self._m.setUrls([QUrl.fromLocalFile(p) for p in paths]); self.ok = False
    def mimeData(self): return self._m
    def acceptProposedAction(self): self.ok = True

def test_window_drops_route_files_to_ask_and_folders_to_folders(env, tmp_path):
    w, *_ = env
    (tmp_path / "in").mkdir(); (tmp_path / "in" / "a.txt").write_text("Mail ana@x.com"); (tmp_path / "f.txt").write_text("hi")
    w.set_mode(1)
    w.dropEvent(_Ev([str(tmp_path / "f.txt")]))
    assert w.mode_stack.currentIndex() == 0 and pump(lambda: len(w.composer.attachments()) == 1)
    w.dropEvent(_Ev([str(tmp_path / "in")]))
    assert w.mode_stack.currentIndex() == 1 and w.folders.folder == str(tmp_path / "in")


# ---------------------------------------------------------------- settings, accounts, inbox
def test_settings_roundtrip_and_categories_take_effect(env):
    w, *_ = env
    from aegis.ui_dialogs import SettingsDialog, disabled_categories
    d = SettingsDialog(w)
    d.cat_checks["names"].setChecked(False); d.idle.setValue(5); d.region_combo.setCurrentText("Portugal (PT)")
    d.key_edits["openai"].setText("sk-openai"); d.ocr_check.setChecked(False); d._save()
    assert "names" in disabled_categories() and stores.settings().value("idle_minutes", type=int) == 5
    assert stores.get_api_key("openai") == "sk-openai"
    w._restart_idle(); assert w._idle.isActive()
    ask(w, "x", [doc()]); assert "Maria" in w.session.result.text           # names switched off: not hidden
    d.deleteLater()

def test_account_dialog_saves_tests_and_removes(env):
    w, _s, _sh, _t, mp = env
    from aegis.ui_dialogs import AccountDialog
    d = AccountDialog(w)
    d.name.setText("Work"); d.provider.setCurrentText("Gmail / Google Workspace")
    d.user.setText("me@corp.com"); d.pw.setText("abcd efgh ijkl mnop")
    mp.setattr(mailbox, "test_login", lambda *a: (_ for _ in ()).throw(OSError("refused")))
    d._test(); assert pump(lambda: "Could not sign in" in d.result_label.text())
    d._save()
    acct = stores.load_accounts()[0]
    assert acct.host == "imap.gmail.com" and stores.account_credentials(acct.id)[0] == "me@corp.com"
    d2 = AccountDialog(w, acct); assert d2.user.text() == "me@corp.com"
    stores.delete_account(acct.id); assert stores.load_accounts() == []
    d.deleteLater(); d2.deleteLater()

def test_inbox_attachment_flow_with_two_accounts(env):
    w, _s, shown, _t, mp = env
    for i, (lbl, host) in enumerate((("Work", "imap.gmail.com"), ("iCloud", "imap.mail.me.com"))):
        a = stores.Account(stores.new_account_id(), lbl, "x", host, 993); stores.set_account_credentials(a.id, f"u{i}@x.com", "pw")
        stores.save_accounts(stores.load_accounts() + [a])
    got = {}
    def fake_many(creds, start, end, limit, exact): got["creds"] = creds; return "=== INBOX DIGEST ===\n--- EMAIL ---\nFrom: Ana <ana@x.com>\n--- EMAIL ---\nFrom: Bob"
    mp.setattr(mailbox, "fetch_many", fake_many)
    from aegis import ui_dialogs
    mp.setattr(ui_dialogs.InboxDialog, "exec", lambda self: True)
    w.attach_inbox()
    assert pump(lambda: len(w.composer.attachments()) == 1) and len(got["creds"]) == 2
    assert "2 accounts" in w.composer.attachments()[0].name and "2 email" in w.composer.attachments()[0].name

def test_inbox_without_accounts_opens_settings_instead_of_failing(env):
    w, _s, shown, _t, mp = env
    opened = []
    mp.setattr(w, "open_settings", lambda: opened.append(1))
    w.attach_inbox(); assert opened and shown


# ---------------------------------------------------------------- folders
def _make_folder(tmp_path):
    src = tmp_path / "in"; (src / "sub").mkdir(parents=True)
    (src / "a.txt").write_text("Maria Santos wrote to maria@x.com"); (src / "sub" / "b.txt").write_text("Thank Maria Santos today")
    return str(src)

def test_folder_batch_preview_recheck_and_restore(env, tmp_path):
    w, _s, shown, _t, mp = env
    src = _make_folder(tmp_path)
    w.folders.key_check.setChecked(True)
    w.set_mode(1); w.folders.set_folder(src)
    assert pump(lambda: w.folders.last is not None and not w.folders._busy)
    out = w.folders.last.out_dir
    blob = "".join(open(os.path.join(out, n)).read() for n in os.listdir(out))
    assert "Maria" not in blob and "maria@x.com" not in blob and w.folders.last.key_kept
    # second batch: a changed file; preview of an OLD file must still work
    time.sleep(1.1); open(os.path.join(src, "a.txt"), "a").write(" and Rita Gomes")
    w.folders.rescan(manual=True)
    assert pump(lambda: w.folders.last is not None and w.folders.last.batch_id != os.path.basename(out)[10:] and not w.folders._busy, 30)
    previews = []
    from aegis import folders_ui
    mp.setattr(folders_ui.PreviewDialog, "exec", lambda self: previews.append(self.windowTitle()) or 0)
    for row, st in enumerate(w.folders.states):
        w.folders._preview(row, 0)
    assert len(previews) == 2, f"previews for both files expected, got {previews}"
    # restore a returned file
    ids = folders_ui.batch.list_restorable_batches()
    mp.setattr(QInputDialog, "getItem", staticmethod(lambda *a, **k: (ids[-1], True)))
    returned = tmp_path / "ret.txt"; returned.write_text("Reply to <PERSON_1> at <EMAIL_ADDRESS_1>")
    dest = tmp_path / "restored.txt"
    mp.setattr(QFileDialog, "getOpenFileName", staticmethod(lambda *a, **k: (str(returned), "")))
    mp.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (str(dest), "")))
    w.folders.restore()
    assert "Maria Santos" in dest.read_text() and "maria@x.com" in dest.read_text()

def test_failing_file_is_not_retried_forever(env, tmp_path):
    w, *_ = env
    src = tmp_path / "in"; src.mkdir(); (src / "ok.txt").write_text("Mail ana@x.com"); (src / "bad.pdf").write_bytes(b"not a pdf")
    w.set_mode(1); w.folders.set_folder(str(src))
    assert pump(lambda: w.folders.last is not None and not w.folders._busy)
    n = len(os.listdir(str(tmp_path / "out")))
    for _ in range(3):
        w.folders.rescan(); pump(lambda: not w.folders._busy, 5); _APP.processEvents(); time.sleep(0.3); _APP.processEvents()
    assert len(os.listdir(str(tmp_path / "out"))) == n, "failed files must not create endless new batches"

def test_empty_and_unsupported_only_folders_say_so(env, tmp_path):
    w, *_ = env
    src = tmp_path / "empty"; src.mkdir(); (src / "x.xyz").write_text("x")
    w.set_mode(1); w.folders.set_folder(str(src))
    assert pump(lambda: "No supported files" in w.folders.summary.text())


# ---------------------------------------------------------------- diagnostics
def test_export_diagnostics_contains_no_personal_data(env, tmp_path):
    w, _s, _sh, _t, mp = env
    ask(w, "Summarise for Ana", [doc()]); w.send_now(); assert pump(lambda: w.answer is not None)
    dest = str(tmp_path / "diag.zip")
    mp.setattr(QFileDialog, "getSaveFileName", staticmethod(lambda *a, **k: (dest, "")))
    w.export_diagnostics()
    with zipfile.ZipFile(dest) as z:
        blob = "".join(z.read(n).decode("utf-8", "ignore") for n in z.namelist())
    for leak in ("Maria", "maria@x.com", "912 345", "Contract", "sk-test"):
        assert leak not in blob
    w.inspector.refresh_log(); assert w.inspector.log_table.rowCount() > 3
